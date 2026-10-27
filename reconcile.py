#!/usr/bin/env python3
"""E-commerce settlement reconciliation engine.

Reads the raw CSV exports in sample_data/ and builds
output/Ecommerce-Profit-Settlement-Reconciliation.xlsx: true profit by SKU,
channel and country, a line-by-line settlement reconciliation, and an
itemised exceptions report.

All money computed with Decimal at order level (rounded to cents) and then
aggregated, so every sheet's totals tie out exactly.
Ad spend is allocated to orders by gross-revenue share within each channel
(largest-remainder method, exact to the cent).
"""
import csv, json
from collections import defaultdict
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = Path(__file__).parent
IN = BASE / "sample_data"
OUT = BASE / "output"
OUT.mkdir(exist_ok=True)
CENT = Decimal("0.01")
def q(x): return Decimal(x).quantize(CENT, rounding=ROUND_HALF_UP)
def D(x): return Decimal(str(x))
def M(d): return float(q(d))          # Decimal -> float for openpyxl
ZERO = Decimal("0")

COST = {"NHG-001": D("18.50"), "NHG-002": D("12.25"), "NHG-003": D("22.00"),
        "NHG-004": D("38.75"), "NHG-005": D("9.80"),  "NHG-006": D("30.40"),
        "NHG-007": D("26.10"), "NHG-008": D("41.60"), "NHG-009": D("14.90"),
        "NHG-010": D("11.35")}

# ------------------------------------------------------------------ load
def read(name):
    with open(IN / name, newline="") as fh:
        return list(csv.DictReader(fh))

orders = read("orders.csv")
amz_settle = read("settlement_amazon.csv")
shp_settle = read("shopify_payouts.csv")
ad_rows = read("ad_spend.csv")
ship = {r["order_id"]: D(r["shipping_cost"]) for r in read("shipping_costs.csv")}

ad_by_channel = defaultdict(lambda: ZERO)
for r in ad_rows:
    ad_by_channel[r["channel"]] += D(r["ad_spend"])

# settlement indexed by order
settle = defaultdict(list)
for r in amz_settle:
    settle[r["order_id"]].append({
        "fees": D(r["referral_fee"]) + D(r["fulfillment_fee"]),
        "refund_ded": D(r["refund_deduction"]), "payout": D(r["payout_amount"]),
        "gross": D(r["gross_sales"]), "ref": "Amazon settlement"})
for r in shp_settle:
    settle[r["order_id"]].append({
        "fees": D(r["processing_fee"]),
        "refund_ded": D(r["refund_deduction"]), "payout": D(r["payout_amount"]),
        "gross": D(r["gross_sales"]), "ref": f"Shopify batch {r['payout_batch']}"})

# ------------------------------------------------------------------ per-order economics
def expected_fees(o, rev):
    if o["channel"] == "Amazon":
        if o["order_status"] == "Refunded":
            return q(rev * D("0.03")) + D("4.25")
        return q(rev * D("0.15")) + D("4.25")
    return q(rev * D("0.029")) + D("0.30")

per = {}   # order_id -> dict of Decimal values
for o in orders:
    rev = q(D(o["unit_price"]) * int(o["quantity"]))
    refunded = o["order_status"] == "Refunded"
    per[o["order_id"]] = {
        "order": o, "rev": rev,
        "refund_rev": rev if refunded else ZERO,
        "cogs": ZERO if refunded else q(COST[o["sku"]] * int(o["quantity"])),
        "fees_exp": expected_fees(o, rev),
        "shipping": ship[o["order_id"]],
        "ad": ZERO,
    }

# allocate ad spend per channel by gross revenue share (largest remainder, exact cents)
for channel, total in ad_by_channel.items():
    ids = [oid for oid, p in per.items() if p["order"]["channel"] == channel]
    chan_rev = sum(per[i]["rev"] for i in ids)
    raw = {i: total * per[i]["rev"] / chan_rev for i in ids}
    floored = {i: (raw[i] * 100).to_integral_value(rounding="ROUND_FLOOR") / 100 for i in ids}
    cents_left = int((total - sum(floored.values())) * 100)
    by_remainder = sorted(ids, key=lambda i: raw[i] - floored[i], reverse=True)
    for i in by_remainder[:cents_left]:
        floored[i] += CENT
    for i in ids:
        per[i]["ad"] = q(floored[i])

for p in per.values():
    p["net"] = q(p["rev"] - p["refund_rev"] - p["cogs"] - p["fees_exp"] - p["ad"] - p["shipping"])

# ------------------------------------------------------------------ reconciliation
recon, exceptions = [], []
for o in orders:
    oid = o["order_id"]; p = per[oid]
    rows = settle.get(oid, [])
    exp_payout = q(p["rev"] - p["fees_exp"] - p["refund_rev"])
    act_payout = q(sum(r["payout"] for r in rows)) if rows else ZERO
    fees_act = q(sum(r["fees"] for r in rows)) if rows else ZERO
    refund_act = q(sum(r["refund_ded"] for r in rows)) if rows else ZERO
    diff = q(act_payout - exp_payout)
    if not rows:
        status, etype = "Missing", "Missing from settlement"
        expl = "Order never appears in the platform settlement — the full expected payout was never received."
    elif len(rows) > 1:
        status, etype = "Duplicate", "Duplicate payout"
        expl = f"Order was paid {len(rows)} times in settlement; the extra payment must be flagged before the platform reclaims it."
    elif refund_act > p["refund_rev"]:
        status, etype = "Underpaid", "Refund deducted twice"
        expl = "The customer refund was deducted twice from the payout (refund deduction exceeds the order revenue)."
    elif fees_act > p["fees_exp"]:
        status, etype = "Underpaid", "Fee overcharge"
        expl = "Platform fees charged are higher than the fee schedule allows."
    elif diff < 0:
        status, etype = "Underpaid", "Wrong payout amount"
        expl = "Payout does not match settlement arithmetic (gross - fees - refunds)."
    elif diff > 0:
        status, etype = "Overpaid", "Overpayment"
        expl = "Payout is higher than the fee schedule and order data support."
    else:
        status, etype = "Matched", ""
        expl = ""
    recon.append({"oid": oid, "channel": o["channel"], "date": o["date"],
                  "gross": p["rev"], "fees_exp": p["fees_exp"], "fees_act": fees_act,
                  "refund_exp": p["refund_rev"], "refund_act": refund_act,
                  "exp": exp_payout, "act": act_payout, "diff": diff,
                  "n": len(rows), "status": status})
    if status != "Matched":
        exceptions.append({"type": etype, "oid": oid, "channel": o["channel"],
                           "exp": exp_payout, "act": act_payout, "impact": diff,
                           "abs": abs(diff), "expl": expl})

# ------------------------------------------------------------------ aggregates
def agg(keyfn):
    out = {}
    for p in per.values():
        k = keyfn(p["order"])
        a = out.setdefault(k, {"units": 0, "rev": ZERO, "refund": ZERO, "cogs": ZERO,
                               "fees": ZERO, "ad": ZERO, "ship": ZERO, "net": ZERO})
        a["units"] += int(p["order"]["quantity"]); a["rev"] += p["rev"]
        a["refund"] += p["refund_rev"]; a["cogs"] += p["cogs"]; a["fees"] += p["fees_exp"]
        a["ad"] += p["ad"]; a["ship"] += p["shipping"]; a["net"] += p["net"]
    return out

by_sku = agg(lambda o: o["sku"])
by_cc = agg(lambda o: (o["channel"], o["country"]))
TOT = {"units": sum(int(o["quantity"]) for o in orders),
       "rev": sum(p["rev"] for p in per.values()), "refund": sum(p["refund_rev"] for p in per.values()),
       "cogs": sum(p["cogs"] for p in per.values()), "fees": sum(p["fees_exp"] for p in per.values()),
       "ad": sum(p["ad"] for p in per.values()), "ship": sum(p["shipping"] for p in per.values()),
       "net": sum(p["net"] for p in per.values())}
EXP_TOT = sum(r["exp"] for r in recon); ACT_TOT = sum(r["act"] for r in recon)
EXC_ABS = sum(e["abs"] for e in exceptions)

# save computed figures for the README step
(OUT / "computed.json").write_text(json.dumps({
    "orders": len(orders), "units": TOT["units"],
    "revenue": str(TOT["rev"]), "refunds": str(TOT["refund"]), "cogs": str(TOT["cogs"]),
    "fees": str(TOT["fees"]), "ad_spend": str(TOT["ad"]), "shipping": str(TOT["ship"]),
    "net_profit": str(TOT["net"]), "expected_payouts": str(EXP_TOT), "actual_payouts": str(ACT_TOT),
    "net_leakage": str(EXP_TOT - ACT_TOT), "exceptions": len(exceptions),
    "exception_abs_impact": str(EXC_ABS),
    "exception_types": {t: sum(1 for e in exceptions if e["type"] == t)
                        for t in sorted({e["type"] for e in exceptions})},
}, indent=2))

# ------------------------------------------------------------------ workbook styling
NAVY = "1F3864"; HDR_FILL = PatternFill("solid", fgColor=NAVY)
HDR_FONT = Font(bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(bold=True, size=14, color=NAVY)
SUB_FONT = Font(bold=True, size=11, color=NAVY)
BOLD = Font(bold=True)
MONEY = '"$"#,##0.00'; MONEY_NEG = '"$"#,##0.00;[Red]-"$"#,##0.00'
THIN = Side(style="thin", color="B0B0B0")
TOT_BORDER = Border(top=Side(style="double", color=NAVY))
STATUS_FILL = {"Matched": PatternFill("solid", fgColor="E2EFDA"),
               "Underpaid": PatternFill("solid", fgColor="FCE4E4"),
               "Overpaid": PatternFill("solid", fgColor="FFF2CC"),
               "Missing": PatternFill("solid", fgColor="F8CBAD"),
               "Duplicate": PatternFill("solid", fgColor="DDEBF7")}

def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HDR_FILL; cell.font = HDR_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = ws.cell(row=row + 1, column=1).coordinate

def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w

wb = Workbook()

# ---- Sheet 1: Profit by SKU
ws = wb.active; ws.title = "Profit by SKU"
ws["A1"] = "Northwind Home Goods — Profit by SKU (October 2026, USD)"; ws["A1"].font = TITLE_FONT
hdr = ["SKU", "Units Sold", "Gross Revenue", "Refunds", "COGS", "Platform Fees (per schedule)",
       "Ad Spend (allocated)", "Shipping", "NET PROFIT"]
ws.append([]); ws.append(hdr); style_header(ws, 3, len(hdr))
for sku in sorted(by_sku):
    a = by_sku[sku]
    ws.append([sku, a["units"], M(a["rev"]), M(a["refund"]), M(a["cogs"]), M(a["fees"]),
               M(a["ad"]), M(a["ship"]), M(a["net"])])
r = ws.max_row + 1
ws.append(["TOTAL", TOT["units"], M(TOT["rev"]), M(TOT["refund"]), M(TOT["cogs"]), M(TOT["fees"]),
           M(TOT["ad"]), M(TOT["ship"]), M(TOT["net"])])
for c in range(1, len(hdr) + 1):
    cell = ws.cell(row=r, column=c); cell.font = BOLD; cell.border = TOT_BORDER
for row in ws.iter_rows(min_row=4, max_row=r, min_col=3, max_col=9):
    for cell in row: cell.number_format = MONEY
widths(ws, {"A": 12, "B": 12, "C": 15, "D": 12, "E": 12, "F": 20, "G": 18, "H": 12, "I": 15})

# ---- Sheet 2: Profit by Channel & Country
ws = wb.create_sheet("Profit by Channel & Country")
ws["A1"] = "Northwind Home Goods — Profit by Channel & Country (October 2026, USD)"; ws["A1"].font = TITLE_FONT
hdr2 = ["Channel", "Country", "Units Sold", "Gross Revenue", "Refunds", "COGS",
        "Platform Fees (per schedule)", "Ad Spend (allocated)", "Shipping", "NET PROFIT"]
ws.append([]); ws.append(hdr2); style_header(ws, 3, len(hdr2))
for (ch, co) in sorted(by_cc):
    a = by_cc[(ch, co)]
    ws.append([ch, co, a["units"], M(a["rev"]), M(a["refund"]), M(a["cogs"]), M(a["fees"]),
               M(a["ad"]), M(a["ship"]), M(a["net"])])
r = ws.max_row + 1
ws.append(["TOTAL", "", TOT["units"], M(TOT["rev"]), M(TOT["refund"]), M(TOT["cogs"]), M(TOT["fees"]),
           M(TOT["ad"]), M(TOT["ship"]), M(TOT["net"])])
for c in range(1, len(hdr2) + 1):
    cell = ws.cell(row=r, column=c); cell.font = BOLD; cell.border = TOT_BORDER
for row in ws.iter_rows(min_row=4, max_row=r, min_col=4, max_col=10):
    for cell in row: cell.number_format = MONEY
widths(ws, {"A": 13, "B": 10, "C": 12, "D": 15, "E": 12, "F": 12, "G": 20, "H": 18, "I": 12, "J": 15})

# ---- Sheet 3: Settlement Reconciliation
ws = wb.create_sheet("Settlement Reconciliation")
ws["A1"] = "Settlement Reconciliation — expected payout (fee schedule) vs actual payout (platform files)"
ws["A1"].font = TITLE_FONT
hdr3 = ["Order ID", "Channel", "Order Date", "Gross Revenue", "Expected Fees", "Actual Fees",
        "Expected Refund Deduction", "Actual Refund Deduction", "Expected Payout", "Actual Payout",
        "Difference (Actual - Expected)", "Settlement Rows", "Status"]
ws.append([]); ws.append(hdr3); style_header(ws, 3, len(hdr3))
for rec in sorted(recon, key=lambda x: x["oid"]):
    ws.append([rec["oid"], rec["channel"], rec["date"], M(rec["gross"]), M(rec["fees_exp"]),
               M(rec["fees_act"]), M(rec["refund_exp"]), M(rec["refund_act"]), M(rec["exp"]),
               M(rec["act"]), M(rec["diff"]), rec["n"], rec["status"]])
    rr = ws.max_row
    ws.cell(row=rr, column=13).fill = STATUS_FILL[rec["status"]]
    ws.cell(row=rr, column=13).font = BOLD
r = ws.max_row + 1
ws.append(["TOTAL", "", "", M(TOT["rev"]), M(TOT["fees"]), M(sum(rr2['fees_act'] for rr2 in recon)),
           M(TOT["refund"]), M(sum(rr2['refund_act'] for rr2 in recon)), M(EXP_TOT), M(ACT_TOT),
           M(ACT_TOT - EXP_TOT), sum(rr2["n"] for rr2 in recon), ""])
for c in range(1, len(hdr3) + 1):
    cell = ws.cell(row=r, column=c); cell.font = BOLD; cell.border = TOT_BORDER
for row in ws.iter_rows(min_row=4, max_row=r, min_col=4, max_col=11):
    for cell in row: cell.number_format = MONEY
ws.auto_filter.ref = f"A3:M{r - 1}"
widths(ws, {"A": 12, "B": 11, "C": 12, "D": 14, "E": 14, "F": 13, "G": 19, "H": 19,
            "I": 15, "J": 14, "K": 21, "L": 14, "M": 12})

# ---- Sheet 4: Exceptions Report
ws = wb.create_sheet("Exceptions Report")
ws["A1"] = "Exceptions Report — every discrepancy found, with its money impact"; ws["A1"].font = TITLE_FONT
hdr4 = ["#", "Exception Type", "Order ID", "Channel", "Expected Payout", "Actual Payout",
        "Money Impact (Actual - Expected)", "Explanation"]
ws.append([]); ws.append(hdr4); style_header(ws, 3, len(hdr4))
for i, e in enumerate(sorted(exceptions, key=lambda x: -x["abs"]), 1):
    ws.append([i, e["type"], e["oid"], e["channel"], M(e["exp"]), M(e["act"]), M(e["impact"]), e["expl"]])
    ws.cell(row=ws.max_row, column=8).alignment = Alignment(wrap_text=True, vertical="top")
r = ws.max_row + 1
ws.append(["", f"TOTAL — {len(exceptions)} exceptions", "", "", M(sum(e['exp'] for e in exceptions)),
           M(sum(e['act'] for e in exceptions)), M(sum(e['impact'] for e in exceptions)),
           f"Gross value of all discrepancies (sum of absolute impacts): ${EXC_ABS:,.2f}"])
for c in range(1, len(hdr4) + 1):
    cell = ws.cell(row=r, column=c); cell.font = BOLD; cell.border = TOT_BORDER
for row in ws.iter_rows(min_row=4, max_row=r, min_col=5, max_col=7):
    for cell in row: cell.number_format = MONEY_NEG
ws.auto_filter.ref = f"A3:H{r - 1}"
widths(ws, {"A": 5, "B": 26, "C": 12, "D": 11, "E": 15, "F": 14, "G": 24, "H": 78})

# ---- Sheet 5: Summary
ws = wb.create_sheet("Summary")
ws["A1"] = "Northwind Home Goods — E-commerce Profit & Settlement Reconciliation"; ws["A1"].font = TITLE_FONT
ws["A2"] = "Sample engagement · October 2026 · Amazon + Shopify · All data 100% fictional"
ws["A2"].font = Font(italic=True, color="595959")
lines = [
    ("", None),
    ("ORDERS & REVENUE", "h"),
    ("Orders processed", len(orders)), ("Units sold", TOT["units"]),
    ("Gross revenue", TOT["rev"]), ("Refunds returned to customers", TOT["refund"]),
    ("Net revenue", TOT["rev"] - TOT["refund"]),
    ("PROFIT", "h"),
    ("Cost of goods sold (COGS)", TOT["cogs"]), ("Platform fees (per fee schedule)", TOT["fees"]),
    ("Advertising spend", TOT["ad"]), ("Shipping costs", TOT["ship"]),
    ("NET PROFIT", TOT["net"]),
    ("SETTLEMENT RECONCILIATION", "h"),
    ("Total expected payouts (fee schedule)", EXP_TOT), ("Total actually paid by platforms", ACT_TOT),
    ("Net leakage (expected - actual)", EXP_TOT - ACT_TOT),
    ("Exceptions found", len(exceptions)),
    ("Gross value of all discrepancies (absolute)", EXC_ABS),
    ("METHOD NOTES", "h"),
    ("Note 1", "Profit uses fees per the client's fee schedule (sample_data/fee_schedule.md); every deviation the platforms actually charged is isolated in the Exceptions Report."),
    ("Note 2", "COGS is counted only on non-refunded orders (refunded goods are assumed returned to inventory). Shipping is charged on all orders, including refunded ones."),
    ("Note 3", "Ad spend is allocated to orders/SKUs by gross-revenue share within each channel, exact to the cent (largest-remainder method)."),
    ("Note 4", "All figures are stored values computed from the raw CSVs in sample_data/ and independently re-verified by validate.py — every total in this workbook ties out exactly across all sheets."),
]
row = 3
for label, val in lines:
    row += 1
    if val == "h":
        ws.cell(row=row, column=1, value=label).font = SUB_FONT
    elif val is None:
        continue
    elif isinstance(val, str):
        ws.cell(row=row, column=1, value=label).font = BOLD
        c = ws.cell(row=row, column=2, value=val); c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row].height = 30
    else:
        ws.cell(row=row, column=1, value=label)
        c = ws.cell(row=row, column=2, value=M(val) if isinstance(val, Decimal) else val)
        if isinstance(val, Decimal): c.number_format = MONEY
        if label in ("NET PROFIT", "Net leakage (expected - actual)"):
            ws.cell(row=row, column=1).font = BOLD; c.font = BOLD
widths(ws, {"A": 44, "B": 95})

out = OUT / "Ecommerce-Profit-Settlement-Reconciliation.xlsx"
wb.save(out)
print(f"workbook saved: {out}")
print(f"revenue={TOT['rev']} net_profit={TOT['net']} exceptions={len(exceptions)} "
      f"abs_impact={EXC_ABS} expected={EXP_TOT} actual={ACT_TOT} leakage={EXP_TOT - ACT_TOT}")
