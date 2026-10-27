#!/usr/bin/env python3
"""Independent validation of the reconciliation workbook.

Recomputes every headline figure from the raw CSVs in sample_data/ with its
own logic (nothing imported from reconcile.py), then reads back
output/Ecommerce-Profit-Settlement-Reconciliation.xlsx and asserts the
numbers match. Prints PASS/FAIL per check.
"""
import csv, sys
from collections import defaultdict
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook

BASE = Path(__file__).parent
IN = BASE / "sample_data"
CENT = Decimal("0.01")
def q(x): return Decimal(x).quantize(CENT, rounding=ROUND_HALF_UP)
def D(x): return Decimal(str(x))
ZERO = Decimal("0")
TOL = Decimal("0.005")

UNIT_COST = {"NHG-001": D("18.50"), "NHG-002": D("12.25"), "NHG-003": D("22.00"),
             "NHG-004": D("38.75"), "NHG-005": D("9.80"), "NHG-006": D("30.40"),
             "NHG-007": D("26.10"), "NHG-008": D("41.60"), "NHG-009": D("14.90"),
             "NHG-010": D("11.35")}

results = []
def check(name, cond, detail=""):
    results.append((name, bool(cond), detail))
    print(f"{'PASS' if cond else 'FAIL'}  {name}" + (f"  [{detail}]" if detail else ""))

def rd(name):
    with open(IN / name, newline="") as fh:
        return list(csv.DictReader(fh))

orders = rd("orders.csv")
ship_map = {r["order_id"]: D(r["shipping_cost"]) for r in rd("shipping_costs.csv")}
ad_total = defaultdict(lambda: ZERO)
for r in rd("ad_spend.csv"):
    ad_total[r["channel"]] += D(r["ad_spend"])

settle_rows = defaultdict(list)
for r in rd("settlement_amazon.csv"):
    settle_rows[r["order_id"]].append(D(r["referral_fee"]) + D(r["fulfillment_fee"]))
    settle_rows[r["order_id"] + "#payout"].append(D(r["payout_amount"]))
    settle_rows[r["order_id"] + "#refund"].append(D(r["refund_deduction"]))
for r in rd("shopify_payouts.csv"):
    settle_rows[r["order_id"]].append(D(r["processing_fee"]))
    settle_rows[r["order_id"] + "#payout"].append(D(r["payout_amount"]))
    settle_rows[r["order_id"] + "#refund"].append(D(r["refund_deduction"]))

def sched_fees(channel, status, rev):
    if channel == "Amazon":
        base = q(rev * D("0.15"))
        return (q(base * D("0.20")) + D("4.25")) if status == "Refunded" else base + D("4.25")
    return q(rev * D("0.029")) + D("0.30")

# per-order recompute
rev_t = ref_t = cogs_t = fees_t = ship_t = ZERO
exp_t = act_t = ZERO
my_exceptions = {}   # order_id -> type
for o in orders:
    rev = q(D(o["unit_price"]) * int(o["quantity"]))
    refunded = o["order_status"] == "Refunded"
    fees = sched_fees(o["channel"], o["order_status"], rev)
    rev_t += rev; fees_t += fees; ship_t += ship_map[o["order_id"]]
    if refunded: ref_t += rev
    else: cogs_t += q(UNIT_COST[o["sku"]] * int(o["quantity"]))
    exp = q(rev - fees - (rev if refunded else ZERO))
    pays = settle_rows.get(o["order_id"] + "#payout", [])
    fee_acts = settle_rows.get(o["order_id"], [])
    ref_acts = settle_rows.get(o["order_id"] + "#refund", [])
    act = sum(pays) if pays else ZERO
    exp_t += exp; act_t += act
    if not pays:
        my_exceptions[o["order_id"]] = "Missing from settlement"
    elif len(pays) > 1:
        my_exceptions[o["order_id"]] = "Duplicate payout"
    elif sum(ref_acts) > (rev if refunded else ZERO):
        my_exceptions[o["order_id"]] = "Refund deducted twice"
    elif sum(fee_acts) > fees:
        my_exceptions[o["order_id"]] = "Fee overcharge"
    elif act != exp:
        my_exceptions[o["order_id"]] = "Wrong payout amount" if act < exp else "Overpayment"

ad_grand = sum(ad_total.values())
net_profit = q(rev_t - ref_t - cogs_t - fees_t - ad_grand - ship_t)

# ------------------------------------------------------------ workbook read
wb = load_workbook(BASE / "output" / "Ecommerce-Profit-Settlement-Reconciliation.xlsx", data_only=True)

def sheet_rows(name):
    ws = wb[name]
    return [[c.value for c in row] for row in ws.iter_rows()]

def find_total(rows, label="TOTAL"):
    for r in rows:
        if r and isinstance(r[0], str) and r[0].strip().upper().startswith(label):
            return r
    return None

def money_eq(a, b):
    return a is not None and abs(D(a) - b) <= TOL

sku = sheet_rows("Profit by SKU")
sku_tot = find_total(sku)
check("SKU sheet totals row present", sku_tot is not None)
if sku_tot:
    check("Gross revenue ties: orders CSV == SKU sheet", money_eq(sku_tot[2], rev_t), f"{rev_t}")
    check("Refunds tie: CSV == SKU sheet", money_eq(sku_tot[3], ref_t), f"{ref_t}")
    check("COGS ties: CSV == SKU sheet", money_eq(sku_tot[4], cogs_t), f"{cogs_t}")
    check("Platform fees tie: CSV == SKU sheet", money_eq(sku_tot[5], fees_t), f"{fees_t}")
    check("Ad spend ties: ad_spend.csv total == SKU sheet", money_eq(sku_tot[6], ad_grand), f"{ad_grand}")
    check("Shipping ties: CSV == SKU sheet", money_eq(sku_tot[7], ship_t), f"{ship_t}")
    check("NET PROFIT ties: recomputed == SKU sheet", money_eq(sku_tot[8], net_profit), f"{net_profit}")
    body = [r for r in sku if r and isinstance(r[0], str) and r[0].startswith("NHG-")]
    colsum = sum(D(r[8]) for r in body)
    check("SKU sheet: sum of SKU net profits == totals row", money_eq(sku_tot[8], colsum), f"{colsum}")
    revsum = sum(D(r[2]) for r in body)
    check("SKU sheet: sum of SKU revenue == totals row", money_eq(sku_tot[2], revsum), f"{revsum}")

cc = sheet_rows("Profit by Channel & Country")
cc_tot = find_total(cc)
check("Channel&Country totals row present", cc_tot is not None)
if cc_tot:
    check("Channel&Country revenue == SKU revenue", money_eq(cc_tot[3], rev_t), f"{cc_tot[3]}")
    check("Channel&Country NET == recomputed net profit", money_eq(cc_tot[9], net_profit), f"{cc_tot[9]}")
    body = [r for r in cc if r and r[0] in ("Amazon", "Shopify")]
    check("Channel&Country: detail rows sum to total net",
          money_eq(cc_tot[9], sum(D(r[9]) for r in body)))

rec = sheet_rows("Settlement Reconciliation")
rec_tot = find_total(rec)
check("Reconciliation totals row present", rec_tot is not None)
if rec_tot:
    check("Recon expected payouts tie", money_eq(rec_tot[8], exp_t), f"{exp_t}")
    check("Recon actual payouts tie", money_eq(rec_tot[9], act_t), f"{act_t}")
    check("Recon difference == actual - expected", money_eq(rec_tot[10], act_t - exp_t))
    body = [r for r in rec if r and isinstance(r[0], str) and r[0].startswith("ORD-")]
    check("Recon covers every order exactly once", len(body) == len(orders), f"{len(body)} rows")
    wb_nonmatched = {r[0]: r[12] for r in body if r[12] != "Matched"}
    check("Recon non-matched order set == independently derived exceptions",
          set(wb_nonmatched) == set(my_exceptions),
          f"wb={len(wb_nonmatched)} mine={len(my_exceptions)}")

exc = sheet_rows("Exceptions Report")
exc_body = [r for r in exc if r and isinstance(r[0], int)]
wb_exc = {r[2]: r[1] for r in exc_body}
check("Exceptions sheet count == derived exceptions", len(exc_body) == len(my_exceptions),
      f"{len(exc_body)} vs {len(my_exceptions)}")
check("Exceptions order IDs match derived set", set(wb_exc) == set(my_exceptions))
type_mismatch = [oid for oid in my_exceptions if wb_exc.get(oid) != my_exceptions[oid]]
check("Exception types match derived types", not type_mismatch, f"mismatch={type_mismatch}")
abs_sum = sum(abs(D(r[6])) for r in exc_body)
signed_sum = sum(D(r[6]) for r in exc_body)
check("Signed exception impacts == actual - expected payouts", money_eq(signed_sum, act_t - exp_t),
      f"{signed_sum} vs {act_t - exp_t}")

summ = sheet_rows("Summary")
flat = {r[0]: r[1] for r in summ if r and isinstance(r[0], str) and len(r) > 1}
check("Summary gross revenue ties", money_eq(flat.get("Gross revenue"), rev_t))
check("Summary NET PROFIT ties", money_eq(flat.get("NET PROFIT"), net_profit))
check("Summary expected payouts tie", money_eq(flat.get("Total expected payouts (fee schedule)"), exp_t))
check("Summary actual payouts tie", money_eq(flat.get("Total actually paid by platforms"), act_t))
check("Summary net leakage == expected - actual",
      money_eq(flat.get("Net leakage (expected - actual)"), exp_t - act_t), f"{exp_t - act_t}")
check("Summary exceptions count ties", flat.get("Exceptions found") == len(my_exceptions),
      f"{flat.get('Exceptions found')}")
check("Summary gross discrepancy value == sum of |impacts| in Exceptions sheet",
      money_eq(flat.get("Gross value of all discrepancies (absolute)"), abs_sum), f"{abs_sum}")

# planted vs found
import json
planted = json.loads((BASE / "sample_data" / "plants.json").read_text())["planted_count"]
check("Planted exceptions == found exceptions", planted == len(my_exceptions),
      f"planted={planted} found={len(my_exceptions)}")

failed = [n for n, ok, _ in results if not ok]
print(f"\n{len(results) - len(failed)}/{len(results)} checks passed")
if failed:
    print("FAILED:", "; ".join(failed)); sys.exit(1)
print("ALL CHECKS PASSED")
